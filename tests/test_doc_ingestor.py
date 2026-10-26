"""
Comprehensive Unit Tests for Phase 4:
- Document Ingestor (DocumentIngestor)
- PDF Page Classification (Digital vs Scanned page OCR branch)
- DOCX, Markdown, Text parsing
- Hierarchical chunking & Breadcrumb prepending
- Long table chunking with repeated header rows
- Qdrant Vector Store indexing & RBAC Hybrid Search
"""

from __future__ import annotations

import io
import os
import sys
from pathlib import Path
import pytest
import fitz  # PyMuPDF
import docx
from PIL import Image

# Add project root to sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from application.services.doc_ingestor import DocumentIngestor, DocumentChunk
from application.interfaces.ocr_engine import OCREngine
from infrastructure.models.bge_m3_embedder import BGEM3Embedder
from infrastructure.vector_store.qdrant_store import QdrantVectorStore
from infrastructure.vector_store.qdrant_config import QdrantSettings
from qdrant_client import QdrantClient


class MockOCREngine(OCREngine):
    """Mock OCR engine returning Vietnamese text and high confidence."""
    def extract_page(self, image_bytes: bytes, lang: str = "vi"):
        return {
            "text": "Nội dung trích xuất từ trang quét bằng OCR tiếng Việt thành công.",
            "confidence": 0.94,
            "tables": ["| Mã SP | Tên SP | Đơn giá |\n|---|---|---|\n| SP01 | Bàn phím | 250.000 |"],
        }


def create_sample_pdf(tmp_path: Path) -> str:
    """Creates a sample PDF with 1 digital page and 1 scanned page (image + minimal text)."""
    pdf_path = str(tmp_path / "sample_policy.pdf")
    doc = fitz.open()

    # Page 1: Digital text page (> 50 chars)
    page1 = doc.new_page()
    text1 = (
        "CONG HOA XA HOI CHU NGHIA VIET NAM\n"
        "Doc lap - Tu do - Hanh phuc\n\n"
        "QUY CHE QUAN LY DU AN CONG NGHE THONG TIN NOI BO\n"
        "Dieu 1. Pham vi dieu chinh va doi tuong ap dung.\n"
        "Quy che nay quy dinh cac nguyen tac, quy trinh va trach nhiem trong viec quan ly du lieu."
    )
    page1.insert_text((50, 72), text1, fontsize=11)

    # Page 2: Scanned page (image inserted, text < 50 chars)
    page2 = doc.new_page()
    # Insert a dummy red image
    img = Image.new("RGB", (200, 200), color=(255, 0, 0))
    img_byte_arr = io.BytesIO()
    img.save(img_byte_arr, format="PNG")
    page2.insert_image(fitz.Rect(50, 50, 250, 250), stream=img_byte_arr.getvalue())
    page2.insert_text((50, 270), "Trang 2", fontsize=10)  # Only 7 chars!

    doc.save(pdf_path)
    doc.close()
    return pdf_path


def create_sample_docx(tmp_path: Path) -> str:
    """Creates a sample DOCX file with headings, paragraphs, and a table."""
    docx_path = str(tmp_path / "sample_manual.docx")
    doc = docx.Document()

    doc.add_heading("Hướng Dẫn Sử Dụng Hệ Thống", level=1)
    doc.add_paragraph("Tài liệu hướng dẫn nhân viên nội bộ về quy trình vận hành và an toàn thông tin.")

    doc.add_heading("Quy Định Bảo Mật", level=2)
    doc.add_paragraph("Mọi nhân viên phải tuân thủ việc phân quyền và bảo mật dữ liệu doanh nghiệp.")

    # Add table
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Cấp độ"
    hdr_cells[1].text = "Quyền truy cập"
    hdr_cells[2].text = "Ghi chú"

    row_cells = table.add_row().cells
    row_cells[0].text = "Level 1"
    row_cells[1].text = "Phòng ban"
    row_cells[2].text = "Chỉ xem dữ liệu nội bộ"

    doc.save(docx_path)
    return docx_path


@pytest.fixture
def in_memory_qdrant_store():
    """Provides an isolated in-memory QdrantVectorStore."""
    client = QdrantClient(location=":memory:")
    settings = QdrantSettings(
        collection_name="test_knowledge",
        dense_dim=1024,
    )
    store = QdrantVectorStore(settings=settings, client=client)
    store.init_collection()
    return store


@pytest.mark.asyncio
async def test_pdf_digital_and_scanned_classification(tmp_path, in_memory_qdrant_store):
    """Verifies that PDF pages are correctly classified into text layer vs OCR branch."""
    pdf_path = create_sample_pdf(tmp_path)
    embedder = BGEM3Embedder()
    ocr_engine = MockOCREngine()

    ingestor = DocumentIngestor(
        embedder=embedder,
        vector_store=in_memory_qdrant_store,
        ocr_engine=ocr_engine,
    )

    res = await ingestor.ingest_file(
        file_path=pdf_path,
        file_id="pdf-file-001",
        department_id="dept-it",
        owner_user_id="user-admin",
        is_public=False,
    )

    assert res.total_pages == 2
    assert res.total_chunks >= 2

    # Chunk from digital page (page 1) should have ocr_confidence == 1.0
    digital_chunks = [c for c in res.chunks if c.page_number == 1]
    assert len(digital_chunks) > 0
    assert digital_chunks[0].ocr_confidence == 1.0
    assert "QUY CHE" in digital_chunks[0].content

    # Chunk from scanned page (page 2) should have OCR confidence from MockOCREngine (0.94)
    scanned_chunks = [c for c in res.chunks if c.page_number == 2]
    assert len(scanned_chunks) > 0
    assert scanned_chunks[0].ocr_confidence == 0.94
    assert "OCR tiếng Việt" in scanned_chunks[0].content


@pytest.mark.asyncio
async def test_docx_and_hierarchical_breadcrumbs(tmp_path, in_memory_qdrant_store):
    """Verifies DOCX parsing, heading hierarchies, breadcrumb generation, and table extraction."""
    docx_path = create_sample_docx(tmp_path)
    embedder = BGEM3Embedder()

    ingestor = DocumentIngestor(
        embedder=embedder,
        vector_store=in_memory_qdrant_store,
    )

    res = await ingestor.ingest_file(
        file_path=docx_path,
        file_id="docx-file-001",
        department_id="dept-hr",
        owner_user_id="user-hr",
        is_public=True,
    )

    assert res.total_chunks >= 2

    # Check breadcrumb prepending
    # Breadcrumb format: "{file_name} > {H1} > {H2}"
    security_chunks = [c for c in res.chunks if "Quy Định Bảo Mật" in c.heading_path]
    assert len(security_chunks) > 0
    assert "Hướng Dẫn Sử Dụng Hệ Thống > Quy Định Bảo Mật" in security_chunks[0].breadcrumb

    # Check Markdown Table extraction from docx
    table_chunks = [c for c in res.chunks if c.chunk_type == "table"]
    assert len(table_chunks) == 1
    assert "| Cấp độ | Quyền truy cập | Ghi chú |" in table_chunks[0].content
    assert "| Level 1 | Phòng ban |" in table_chunks[0].content


@pytest.mark.asyncio
async def test_table_chunking_with_repeated_header(tmp_path, in_memory_qdrant_store):
    """Verifies that large tables exceeding token limits are split and repeat header rows."""
    embedder = BGEM3Embedder()
    # Set small max_chunk_tokens to force table split
    ingestor = DocumentIngestor(
        embedder=embedder,
        vector_store=in_memory_qdrant_store,
        max_chunk_tokens=30,  # small token limit
    )

    # Markdown table with 10 rows
    header = "| Mã NV | Họ và tên | Chức danh |\n|---|---|---|\n"
    rows = "\n".join([f"| NV{i:02d} | Nguyễn Văn {chr(65+i)} | Chuyên viên {i} |" for i in range(10)])
    long_table = header + rows

    chunks = ingestor._chunk_table(
        table_content=long_table,
        doc_id="doc-test",
        file_id="file-test",
        file_name="staff.md",
        breadcrumb="staff.md > Nhân sự",
        heading_path="Nhân sự",
        section_title="Bảng Nhân sự",
        page_num=1,
        owner_user_id="u1",
        department_id="d1",
        is_public=False,
        confidence=1.0,
    )

    # Must be split into multiple chunks
    assert len(chunks) > 1

    # Every split table chunk MUST contain the header and delimiter!
    for ch in chunks:
        assert "| Mã NV | Họ và tên | Chức danh |" in ch.content
        assert "|---|---|---|" in ch.content
        assert ch.chunk_type == "table"


@pytest.mark.asyncio
async def test_qdrant_rbac_hybrid_search(tmp_path, in_memory_qdrant_store):
    """Verifies Qdrant VectorStore end-to-end upsert, RBAC filtering, and delete_by_file_id."""
    embedder = BGEM3Embedder()

    # Document 1: Private to dept-kd
    chunk1 = {
        "chunk_id": "c1111111-1111-1111-1111-111111111111",
        "dense_vector": embedder.embed_query("Quy định hoa hồng kinh doanh")["dense"],
        "sparse_indices": [10, 20, 30],
        "sparse_values": [1.0, 2.0, 1.5],
        "payload": {
            "file_id": "f-kd-01",
            "department_id": "dept-kd",
            "is_public": False,
            "owner_user_id": "user-kd-01",
            "content": "Chính sách hoa hồng phòng kinh doanh năm 2024.",
        },
    }

    # Document 2: Public company-wide
    chunk2 = {
        "chunk_id": "c2222222-2222-2222-2222-222222222222",
        "dense_vector": embedder.embed_query("Nội quy lao động chung")["dense"],
        "sparse_indices": [10, 40, 50],
        "sparse_values": [1.2, 1.8, 2.2],
        "payload": {
            "file_id": "f-pub-01",
            "department_id": "dept-hr",
            "is_public": True,
            "owner_user_id": "user-hr-01",
            "content": "Nội quy lao động áp dụng cho toàn thể cán bộ nhân viên.",
        },
    }

    in_memory_qdrant_store.upsert_chunks([chunk1, chunk2])

    # Search Query
    q_emb = embedder.embed_query("hoa hồng kinh doanh")

    # Case A: User from dept-kd -> should find chunk1 (own dept) and chunk2 (public)
    hits_kd = in_memory_qdrant_store.hybrid_search(
        query_dense=q_emb["dense"],
        query_sparse_indices=q_emb["sparse_indices"],
        query_sparse_values=q_emb["sparse_values"],
        department_id="dept-kd",
        user_id="user-kd-02",
        is_admin=False,
    )
    hit_ids_kd = [h["chunk_id"] for h in hits_kd]
    assert "c1111111-1111-1111-1111-111111111111" in hit_ids_kd
    assert "c2222222-2222-2222-2222-222222222222" in hit_ids_kd

    # Case B: User from dept-it -> should NOT find chunk1 (private dept-kd), but CAN find chunk2 (public)
    hits_it = in_memory_qdrant_store.hybrid_search(
        query_dense=q_emb["dense"],
        query_sparse_indices=q_emb["sparse_indices"],
        query_sparse_values=q_emb["sparse_values"],
        department_id="dept-it",
        user_id="user-it-01",
        is_admin=False,
    )
    hit_ids_it = [h["chunk_id"] for h in hits_it]
    assert "c1111111-1111-1111-1111-111111111111" not in hit_ids_it
    assert "c2222222-2222-2222-2222-222222222222" in hit_ids_it

    # Case C: Admin -> sees everything
    hits_admin = in_memory_qdrant_store.hybrid_search(
        query_dense=q_emb["dense"],
        query_sparse_indices=q_emb["sparse_indices"],
        query_sparse_values=q_emb["sparse_values"],
        department_id="dept-it",
        user_id="user-admin",
        is_admin=True,
    )
    assert len(hits_admin) == 2

    # Case D: Delete by file_id
    del_ok = in_memory_qdrant_store.delete_by_file_id("f-pub-01")
    assert del_ok is True
    hits_after_del = in_memory_qdrant_store.hybrid_search(
        query_dense=q_emb["dense"],
        query_sparse_indices=q_emb["sparse_indices"],
        query_sparse_values=q_emb["sparse_values"],
        department_id="dept-kd",
        user_id="user-kd-01",
        is_admin=True,
    )
    assert "c2222222-2222-2222-2222-222222222222" not in [h["chunk_id"] for h in hits_after_del]
