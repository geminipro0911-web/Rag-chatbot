"""
Comprehensive Unit Tests for Phase 3:
- Tabular Ingestion (ExcelPreprocessor)
- Multi-tier headers ('Cha - Con')
- Merged cells rectangular unmerge & fill
- VN Numbers vs International Numbers, Currency & Percentages
- VN Dates and Excel Serial Numbers (origin 1899-12-30)
- Duplicate column names & alias generation
- Summary rows removal ('Tổng cộng', 'Cộng', 'Total')
- Parquet export with pyarrow zstd
- CatalogService & Synthetic Few-Shot validation with DuckDB sandbox
"""

from __future__ import annotations

import os
import sys
from datetime import datetime
from pathlib import Path
import tempfile
import pytest
import openpyxl
import pandas as pd
import duckdb

# Add project root to sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from application.services.excel_preprocessor import ExcelPreprocessor
from application.services.catalog_service import CatalogService, FewShotPair, SyntheticFewShotOutput
from application.interfaces.llm_client import LLMClient


def create_sample_workbook(tmp_path: str) -> str:
    """Creates a complex multi-feature workbook for testing."""
    wb = openpyxl.Workbook()

    # Sheet 1: Complex sales report with 2-tier header, merged cells, VN numbers, dates, summary row
    ws = wb.active
    ws.title = "Báo Cáo Doanh Thu"

    # Row 1: Title (should be skipped by header detector)
    ws.cell(row=1, column=1, value="CÔNG TY CỔ PHẦN CÔNG NGHỆ ABC - BÁO CÁO TÀI CHÍNH")

    # Row 2 & 3: 2-tier header with merged cells
    # Column A: "Khu vực" (merged A2:A3)
    ws.cell(row=2, column=1, value="Khu vực")
    ws.merge_cells("A2:A3")

    # Column B: "Ngày ghi nhận" (merged B2:B3)
    ws.cell(row=2, column=2, value="Ngày ghi nhận")
    ws.merge_cells("B2:B3")

    # Column C & D: "Doanh thu" parent header (merged C2:D2)
    ws.cell(row=2, column=3, value="Doanh thu")
    ws.cell(row=2, column=4, value="Doanh thu")
    ws.merge_cells("C2:D2")
    ws.cell(row=3, column=3, value="Kế hoạch")
    ws.cell(row=3, column=4, value="Thực hiện")

    # Column E: "Tỷ lệ đạt"
    ws.cell(row=2, column=5, value="Tỷ lệ đạt")
    ws.cell(row=3, column=5, value="Tỷ lệ đạt")

    # Column F: Duplicate column name for testing deduplication ("Doanh thu")
    ws.cell(row=2, column=6, value="Doanh thu")
    ws.cell(row=3, column=6, value="Dự báo")

    # Data Rows (Row 4, 5, 6)
    # Row 4
    ws.cell(row=4, column=1, value="Miền Bắc")
    ws.cell(row=4, column=2, value="ngày 15 tháng 03 năm 2024")
    ws.cell(row=4, column=3, value="1.500.000.000 VNĐ")
    ws.cell(row=4, column=4, value="1.650.000.000 đ")
    ws.cell(row=4, column=5, value="110%")
    ws.cell(row=4, column=6, value="1.700.000.000")

    # Row 5 (merged cells across region A5:A6)
    ws.cell(row=5, column=1, value="Miền Nam")
    ws.merge_cells("A5:A6")
    ws.cell(row=5, column=2, value="20/04/2024")
    ws.cell(row=5, column=3, value="2.000.000.000")
    ws.cell(row=5, column=4, value="1.950.000.000")
    ws.cell(row=5, column=5, value="97,5%")
    ws.cell(row=5, column=6, value="2.100.000.000")

    # Row 6
    ws.cell(row=6, column=2, value="21/04/2024")
    ws.cell(row=6, column=3, value="800.000.000")
    ws.cell(row=6, column=4, value="850.000.000")
    ws.cell(row=6, column=5, value="106,25%")
    ws.cell(row=6, column=6, value="900.000.000")

    # Row 7: Summary row that MUST be stripped
    ws.cell(row=7, column=1, value="Tổng cộng")
    ws.cell(row=7, column=3, value="4.300.000.000")
    ws.cell(row=7, column=4, value="4.450.000.000")
    ws.cell(row=7, column=5, value="103,5%")

    # Sheet 2: Excel Serial Dates & International Numbers
    ws2 = wb.create_sheet(title="Dự Án")
    ws2.append(["Mã dự án", "Tên dự án", "Hạn nộp báo cáo", "Kinh phí USD"])
    # 45383 corresponds to 2024-04-01 in Excel serial numbers (origin 1899-12-30)
    # 45413 corresponds to 2024-05-01
    ws2.append(["PRJ01", "Hệ thống AI", 45383, "$1,250,500.50"])
    ws2.append(["PRJ02", "Nâng cấp Hạ tầng", 45413, "$850,000.00"])

    # Sheet 3: Hidden sheet (should be skipped)
    ws3 = wb.create_sheet(title="Sheet_Bi_An")
    ws3.append(["Dữ liệu ẩn"])
    ws3.sheet_state = "hidden"

    file_path = os.path.join(tmp_path, "test_report.xlsx")
    wb.save(file_path)
    return file_path


def test_excel_preprocessor_end_to_end(tmp_path):
    excel_path = create_sample_workbook(str(tmp_path))
    parquet_out_dir = str(tmp_path / "parquet_output")

    preprocessor = ExcelPreprocessor(output_dir=parquet_out_dir)
    res = preprocessor.process_file(
        file_path=excel_path,
        file_id="file-uuid-001",
        department_id="dept-kd",
        owner_user_id="user-001",
        is_public=False,
        skip_hidden=True,
    )

    # 1. Verification of Sheet Count: 2 visible sheets processed, hidden sheet skipped
    assert len(res.sheets) == 2
    sheet_names = [s.sheet_name for s in res.sheets]
    assert "Báo Cáo Doanh Thu" in sheet_names
    assert "Dự Án" in sheet_names
    assert "Sheet_Bi_An" not in sheet_names

    # Check Sheet 1: "Báo Cáo Doanh Thu"
    s1 = next(s for s in res.sheets if s.sheet_name == "Báo Cáo Doanh Thu")

    # 2. Merged cell rectangular fill verification
    # A5:A6 was merged with "Miền Nam". Both rows 5 and 6 must have "Miền Nam"
    df1 = s1.dataframe
    assert len(df1) == 3  # 3 data rows (rows 4, 5, 6). Row 7 ("Tổng cộng") was stripped!
    assert df1["khu_vuc"].tolist() == ["Miền Bắc", "Miền Nam", "Miền Nam"]

    # 3. 2-tier header combination verification
    # Columns should be combined:
    # "Doanh thu" + "Kế hoạch" -> "Doanh thu - Kế hoạch"
    # "Doanh thu" + "Thực hiện" -> "Doanh thu - Thực hiện"
    # "Doanh thu" + "Dự báo" -> "Doanh thu - Dự báo"
    col_aliases = [c.column_alias for c in s1.columns_schema]
    orig_names = [c.original_name for c in s1.columns_schema]

    assert "Doanh thu - Kế hoạch" in orig_names
    assert "Doanh thu - Thực hiện" in orig_names
    assert "Doanh thu - Dự báo" in orig_names

    assert "doanh_thu_ke_hoach" in col_aliases
    assert "doanh_thu_thuc_hien" in col_aliases
    assert "doanh_thu_du_bao" in col_aliases

    # 4. Vietnamese Date parsing verification
    # Row 4 had "ngày 15 tháng 03 năm 2024"
    # Row 5 had "20/04/2024"
    date_col = s1.columns_schema[1]
    assert date_col.data_type == "date"
    first_date = pd.to_datetime(df1["ngay_ghi_nhan"].iloc[0])
    assert first_date.year == 2024
    assert first_date.month == 3
    assert first_date.day == 15

    # 5. Vietnamese currency & number normalization
    # "1.500.000.000 VNĐ" -> float 1500000000.0
    ke_hoach_col = s1.columns_schema[2]
    assert ke_hoach_col.data_type in ("integer", "float")
    assert df1["doanh_thu_ke_hoach"].iloc[0] == 1500000000.0

    # 6. Percentage normalization
    # "110%" -> 1.10; "97,5%" -> 0.975
    pct_col = s1.columns_schema[4]
    assert pct_col.data_type == "float"
    assert round(df1["ty_le_dat"].iloc[0], 2) == 1.10
    assert round(df1["ty_le_dat"].iloc[1], 3) == 0.975

    # 7. Summary row removal verification
    assert not any("tổng" in str(v).lower() for v in df1["khu_vuc"])
    assert any("Removed summary row" in w for w in s1.warnings)

    # 8. Parquet file verification with DuckDB
    assert os.path.exists(s1.parquet_path)
    con = duckdb.connect()
    res_duck = con.execute(f"SELECT COUNT(*) FROM read_parquet('{s1.parquet_path}')").fetchone()
    assert res_duck[0] == 3
    con.close()

    # 9. Sheet 2 Verification: Excel Serial Date & International Numbers
    s2 = next(s for s in res.sheets if s.sheet_name == "Dự Án")
    df2 = s2.dataframe

    # Serial number 45383 converted to 2024-04-01
    date_serial_col = s2.columns_schema[2]
    assert date_serial_col.data_type == "date"
    d1 = pd.to_datetime(df2["han_nop_bao_cao"].iloc[0])
    assert d1.year == 2024
    assert d1.month == 4
    assert d1.day == 1

    # International currency "$1,250,500.50" -> 1250500.5
    usd_col = s2.columns_schema[3]
    assert usd_col.data_type == "float"
    assert df2["kinh_phi_usd"].iloc[0] == 1250500.50


@pytest.mark.asyncio
async def test_catalog_service_synthetic_fewshot_and_duckdb_sandbox(tmp_path):
    """Tests CatalogService DuckDB sandbox verification for synthetic queries."""
    excel_path = create_sample_workbook(str(tmp_path))
    parquet_out_dir = str(tmp_path / "parquet_output")

    preprocessor = ExcelPreprocessor(output_dir=parquet_out_dir)
    res = preprocessor.process_file(
        file_path=excel_path,
        file_id="file-uuid-002",
        department_id="dept-kd",
        owner_user_id="user-001",
    )
    s1 = res.sheets[0]

    # Mock LLM Client that generates 3 queries:
    # 1. Valid query returning rows (should pass)
    # 2. Syntax error query (should fail)
    # 3. Valid query that returns 0 rows (should fail)
    class MockLLM(LLMClient):
        async def chat_completion(self, messages, temperature=0.7, max_tokens=None):
            return ""

        async def structured_output(self, messages, response_model, temperature=0.2, max_retries=3):
            return SyntheticFewShotOutput(
                examples=[
                    FewShotPair(
                        question="Tổng doanh thu thực hiện theo kế hoạch là bao nhiêu?",
                        sql_query=f"SELECT SUM(doanh_thu_thuc_hien) AS tong_thuc_hien FROM {s1.table_id} WHERE khu_vuc = 'Miền Bắc';",
                        difficulty="easy"
                    ),
                    FewShotPair(
                        question="Lỗi cú pháp SQL?",
                        sql_query="SELEC * FROM NOT_EXISTING_TABLE",
                        difficulty="medium"
                    ),
                    FewShotPair(
                        question="Tìm khu vực Miền Tây?",
                        sql_query=f"SELECT * FROM {s1.table_id} WHERE khu_vuc = 'Miền Tây';",
                        difficulty="easy"
                    )
                ]
            )

    catalog_svc = CatalogService(llm_client=MockLLM(), enable_synthetic_fewshot=True)

    # Test Sandbox Validation Directly
    # Valid query returning data
    valid_sql = f"SELECT * FROM {s1.table_id} WHERE khu_vuc = 'Miền Bắc';"
    assert catalog_svc.validate_sql_query(valid_sql, s1.parquet_path, s1.table_id) is True

    # Invalid SQL syntax
    assert catalog_svc.validate_sql_query("SELEC BROKEN SQL", s1.parquet_path, s1.table_id) is False

    # Valid SQL but returns 0 rows
    zero_rows_sql = f"SELECT * FROM {s1.table_id} WHERE khu_vuc = 'Không Tồn Tại';"
    assert catalog_svc.validate_sql_query(zero_rows_sql, s1.parquet_path, s1.table_id) is False

    # Run full generate_and_validate_fewshots pipeline
    validated_fewshots = await catalog_svc.generate_and_validate_fewshots(s1)

    # Only query 1 should pass! Queries 2 and 3 must be rejected by sandbox.
    assert len(validated_fewshots) == 1
    assert "Tổng doanh thu" in validated_fewshots[0]["question"]
    assert validated_fewshots[0]["source"] == "synthetic"

    # Test ORM Entity generation
    desc = "Báo cáo doanh thu tài chính theo vùng miền."
    keywords = ["doanh thu", "kế hoạch", "thực hiện"]
    table_orm, cols_orm, vals_orm, fs_orm = catalog_svc.build_orm_entities(
        processed_sheet=s1,
        description=desc,
        keywords=keywords,
        fewshots=validated_fewshots,
    )

    assert table_orm.sheet_name == "Báo Cáo Doanh Thu"
    assert table_orm.row_count == 3
    assert len(cols_orm) == len(s1.columns_schema)
    assert len(fs_orm) == 1
    assert fs_orm[0].sql_query == validated_fewshots[0]["sql_query"]


def test_unmerge_rectangular_fill(tmp_path):
    """Verifies that rectangular merged cells fill top-left across all rows and columns."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MergedTest"

    # Set top-left value at B2
    ws.cell(row=2, column=2, value="BlockValue")
    # Merge B2:D4 (3x3 rectangular block)
    ws.merge_cells("B2:D4")

    preprocessor = ExcelPreprocessor()
    preprocessor._unmerge_and_fill(ws)

    # Assert that all 9 cells in B2:D4 contain 'BlockValue'
    for r in range(2, 5):
        for c in range(2, 5):
            assert ws.cell(row=r, column=c).value == "BlockValue", f"Cell ({r},{c}) was not filled!"


def test_column_deduplication_and_aliases():
    """Verifies duplicate column names and SQL-friendly unidecoded snake_case alias generation."""
    preprocessor = ExcelPreprocessor()
    raw_headers = [
        "Doanh thu",
        "Doanh thu",
        "Doanh thu",
        "Số lượng (SP)",
        "Ngày bắt đầu",
        "123 Thử nghiệm",
        "",
        "Unnamed: 7",
    ]

    cleaned_names, aliases = preprocessor._clean_and_alias_columns(raw_headers)

    # Deduplication in display names
    assert cleaned_names[0] == "Doanh thu"
    assert cleaned_names[1] == "Doanh thu_2"
    assert cleaned_names[2] == "Doanh thu_3"
    assert cleaned_names[3] == "Số lượng (SP)"
    assert cleaned_names[4] == "Ngày bắt đầu"
    assert cleaned_names[6] == "col_7"
    assert cleaned_names[7] == "col_8"

    # Aliases must be clean SQL identifiers
    assert aliases[0] == "doanh_thu"
    assert aliases[1] == "doanh_thu_2"
    assert aliases[2] == "doanh_thu_3"
    assert aliases[3] == "so_luong_sp"
    assert aliases[4] == "ngay_bat_dau"
    assert aliases[5] == "col_123_thu_nghiem"  # Prepends 'col_' if starts with digit
    assert aliases[6] == "col_7"


def test_vietnamese_number_and_percentage_parsing():
    """Verifies Vietnamese number formatting, currency stripping, and percentages."""
    preprocessor = ExcelPreprocessor()

    # Column with VN currency
    series_currency = pd.Series(["1.500.000 VNĐ", "2.350.000 đ", "800.000₫", None])
    res_num, is_num, num_type = preprocessor._try_parse_number_column("doanh_thu", series_currency, 3)
    assert is_num is True
    assert res_num.iloc[0] == 1500000.0
    assert res_num.iloc[1] == 2350000.0
    assert res_num.iloc[2] == 800000.0

    # Column with VN decimal and percentages
    series_pct = pd.Series(["15,5%", "84,5%", "100%", None])
    res_pct, is_pct, pct_type = preprocessor._try_parse_number_column("ty_le", series_pct, 3)
    assert is_pct is True
    assert pct_type == "float"
    assert round(res_pct.iloc[0], 4) == 0.1550
    assert round(res_pct.iloc[1], 4) == 0.8450
    assert round(res_pct.iloc[2], 2) == 1.0


def test_summary_rows_stripping_explicit():
    """Verifies that 'Tổng', 'Tổng cộng', 'Cộng' rows in first columns are cleanly stripped."""
    preprocessor = ExcelPreprocessor()
    df = pd.DataFrame({
        "chi_nhanh": ["Hà Nội", "Đà Nẵng", "TP.HCM", "Tổng cộng", "Ghi chú"],
        "so_tien": [100, 200, 300, 600, 0],
    })

    cleaned_df, warnings = preprocessor._strip_summary_rows(df)
    assert len(cleaned_df) == 4
    assert "Tổng cộng" not in cleaned_df["chi_nhanh"].tolist()
    assert len(warnings) == 1
    assert "Removed summary row" in warnings[0]


def test_csv_preprocessor_pipeline(tmp_path):
    """Verifies end-to-end processing of CSV files including type normalization and schema extraction."""
    csv_file = tmp_path / "sales_report.csv"
    csv_content = (
        "Mã Sản Phẩm,Tên Sản Phẩm,Doanh Thu,Ngày Bán\n"
        "SP01,Máy tính xách tay,15.000.000 VNĐ,15/03/2024\n"
        "SP02,Chuột không dây,500.000 VNĐ,16/03/2024\n"
        "Tổng cộng,,15.500.000 VNĐ,\n"
    )
    csv_file.write_text(csv_content, encoding="utf-8")

    out_dir = str(tmp_path / "parquet_out")
    preprocessor = ExcelPreprocessor(output_dir=out_dir)

    result = preprocessor.process_file(
        file_path=str(csv_file),
        file_id="csv-file-001",
        department_id="dep_sales",
        owner_user_id="user_sales",
        is_public=False,
    )

    assert result.file_id == "csv-file-001"
    assert len(result.sheets) == 1

    sheet = result.sheets[0]
    assert sheet.row_count == 2  # Summary row 'Tổng cộng' stripped
    assert os.path.exists(sheet.parquet_path)
    assert len(sheet.columns_schema) == 4

    col_names = [col.column_alias for col in sheet.columns_schema]
    assert "ma_san_pham" in col_names
    assert "doanh_thu" in col_names

    # Check Parquet content in DuckDB
    con = duckdb.connect()
    res = con.execute(f"SELECT COUNT(*), SUM(doanh_thu) FROM read_parquet('{sheet.parquet_path}')").fetchall()
    assert res[0][0] == 2
    assert res[0][1] == 15500000.0

