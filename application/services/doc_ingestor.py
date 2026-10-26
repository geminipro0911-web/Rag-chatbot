"""
Document Ingestion Pipeline Service.
Handles enterprise-grade unstructured document ingestion (.pdf, .docx, .md, .txt).
Follows Module 1 Nhánh B specifications:
- PDF page classification (scanned page vs text layer).
- OCR engine integration for scanned pages (300 DPI render, ocr_confidence tracking).
- DOCX paragraph, heading, and table extraction into Markdown.
- Markdown and text parsing.
- Hierarchical chunking (H1 > H2 breadcrumb prepending, max 512 tokens, 64 overlap).
- Table chunking with header repetition when table exceeds token limit.
- Complete chunk metadata tracking (chunk_id, doc_id, file_id, file_name, page_number,
  section_title, heading_path, chunk_type, owner_user_id, department_id, is_public, ocr_confidence).
- BGE-M3 Dense (1024d) + Sparse lexical indexing into Qdrant.
"""

from __future__ import annotations

import asyncio
import io
import logging
import os
import re
import uuid
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import fitz  # PyMuPDF
import docx

from application.interfaces.embedder import Embedder
from application.interfaces.ocr_engine import OCREngine
from application.interfaces.vector_store import VectorStore
from infrastructure.database.models import DocumentModel

logger = logging.getLogger(__name__)


@dataclass
class DocumentChunk:
    chunk_id: str
    doc_id: str
    file_id: str
    file_name: str
    content: str
    breadcrumb: str
    heading_path: str
    section_title: str
    page_number: int
    chunk_type: str  # "text" or "table"
    owner_user_id: str
    department_id: str
    is_public: bool
    ocr_confidence: float = 1.0


@dataclass
class DocumentIngestResult:
    doc_id: str
    file_id: str
    title: str
    total_pages: int
    total_chunks: int
    chunks: List[DocumentChunk]
    warnings: List[str] = field(default_factory=list)


class DocumentIngestor:
    """Ingests unstructured documents (.pdf, .docx, .md, .txt), chunks them, and indexes into Qdrant."""

    def __init__(
        self,
        embedder: Embedder,
        vector_store: VectorStore,
        ocr_engine: Optional[OCREngine] = None,
        max_chunk_tokens: int = 512,
        chunk_overlap_tokens: int = 64,
    ):
        self.embedder = embedder
        self.vector_store = vector_store
        self.ocr_engine = ocr_engine
        self.max_chunk_tokens = max_chunk_tokens
        self.chunk_overlap_tokens = chunk_overlap_tokens

    async def ingest_file(
        self,
        file_path: str,
        file_id: str,
        department_id: str,
        owner_user_id: str,
        is_public: bool = False,
    ) -> DocumentIngestResult:
        """Processes and indexes a document file into vector database and returns metadata."""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        raw_basename = os.path.basename(file_path)
        file_name = raw_basename[len(file_id) + 1:] if file_id and raw_basename.startswith(f"{file_id}_") else raw_basename
        ext = Path(file_path).suffix.lower()
        doc_id = str(uuid.uuid4())
        warnings: List[str] = []

        # 1. Parse raw sections by file type (offloaded to threadpool to avoid event loop blocking)
        if ext == ".pdf":
            sections, total_pages, parse_warnings = await asyncio.to_thread(self._parse_pdf, file_path)
        elif ext in (".docx", ".doc"):
            sections, total_pages, parse_warnings = await asyncio.to_thread(self._parse_docx, file_path)
        elif ext == ".md":
            sections, total_pages, parse_warnings = await asyncio.to_thread(self._parse_markdown, file_path)
        elif ext == ".txt":
            sections, total_pages, parse_warnings = await asyncio.to_thread(self._parse_txt, file_path)
        else:
            raise ValueError(f"Unsupported document format: {ext}")

        warnings.extend(parse_warnings)

        # 2. Hierarchical chunking
        chunks = self._hierarchical_chunk(
            sections=sections,
            doc_id=doc_id,
            file_id=file_id,
            file_name=file_name,
            owner_user_id=owner_user_id,
            department_id=department_id,
            is_public=is_public,
        )

        if not chunks:
            msg = f"Document '{file_name}' yielded 0 extractable chunks."
            logger.warning(msg)
            warnings.append(msg)
            return DocumentIngestResult(
                doc_id=doc_id,
                file_id=file_id,
                title=file_name,
                total_pages=total_pages,
                total_chunks=0,
                chunks=[],
                warnings=warnings,
            )

        # 3. Embedding (Dense 1024d + Sparse lexical)
        # Note: Embedding text includes prepended breadcrumb: "{file_name} > {H1} > {H2}\n{content}"
        texts_to_embed = [
            f"{c.breadcrumb}\n{c.content}" if c.breadcrumb else c.content
            for c in chunks
        ]
        embedding_results = await asyncio.to_thread(self.embedder.embed_documents, texts_to_embed)

        # 4. Prepare batch for Qdrant
        points_to_upsert: List[Dict[str, Any]] = []
        for chunk_idx, (chunk, emb) in enumerate(zip(chunks, embedding_results)):
            payload = {
                "chunk_id": chunk.chunk_id,
                "chunk_index": chunk_idx,
                "doc_id": chunk.doc_id,
                "file_id": chunk.file_id,
                "file_name": chunk.file_name,
                "page_number": chunk.page_number,
                "page": chunk.page_number,
                "section_title": chunk.section_title,
                "heading_path": chunk.heading_path,
                "chunk_type": chunk.chunk_type,
                "owner_user_id": chunk.owner_user_id,
                "department_id": chunk.department_id,
                "is_public": chunk.is_public,
                "ocr_confidence": chunk.ocr_confidence,
                "content": chunk.content,
                "text": chunk.content,
                "breadcrumb": chunk.breadcrumb,
            }
            points_to_upsert.append({
                "chunk_id": chunk.chunk_id,
                "dense_vector": emb["dense"],
                "sparse_indices": emb["sparse_indices"],
                "sparse_values": emb["sparse_values"],
                "payload": payload,
            })

        # 5. Upsert into Qdrant
        await asyncio.to_thread(self.vector_store.upsert_chunks, points_to_upsert)

        return DocumentIngestResult(
            doc_id=doc_id,
            file_id=file_id,
            title=file_name,
            total_pages=total_pages,
            total_chunks=len(chunks),
            chunks=chunks,
            warnings=warnings,
        )

    # -------------------------------------------------------------------------
    # PDF Parsing & Scanned Page Detection
    # -------------------------------------------------------------------------
    def _parse_pdf(self, file_path: str) -> Tuple[List[Dict[str, Any]], int, List[str]]:
        """
        Parses PDF. Classifies each page:
        - len(text) < 50 AND len(images) > 0 -> Scanned page -> OCR branch.
        - Otherwise -> Digital page text layer.
        """
        sections: List[Dict[str, Any]] = []
        warnings: List[str] = []

        with fitz.open(file_path) as doc:
            total_pages = len(doc)

            for page_idx in range(total_pages):
                page = doc[page_idx]
                page_num = page_idx + 1
                raw_text = page.get_text().strip()
                images = page.get_images()

                is_scanned = (len(raw_text) < 50) and (len(images) > 0)

                if is_scanned:
                    if self.ocr_engine is not None:
                        try:
                            pix = page.get_pixmap(dpi=300)
                            img_bytes = pix.tobytes("png")
                            ocr_res = self.ocr_engine.extract_page(img_bytes, lang="vi")
                            ocr_text = ocr_res.get("text", "")
                            confidence = float(ocr_res.get("confidence", 0.85))
                            ocr_tables = ocr_res.get("tables", [])

                            sections.append({
                                "type": "text",
                                "content": ocr_text,
                                "page": page_num,
                                "heading": f"Trang {page_num}",
                                "confidence": confidence,
                            })
                            for tbl in ocr_tables:
                                sections.append({
                                    "type": "table",
                                    "content": tbl,
                                    "page": page_num,
                                    "heading": f"Bảng - Trang {page_num}",
                                    "confidence": confidence,
                                })
                            continue
                        except Exception as exc:
                            warnings.append(f"OCR failed for page {page_num}: {exc}")

                    # If no OCR engine or OCR failed, keep whatever raw text exists
                    sections.append({
                        "type": "text",
                        "content": raw_text or f"[Trang quét không có text layer - Trang {page_num}]",
                        "page": page_num,
                        "heading": f"Trang {page_num}",
                        "confidence": 0.5,
                    })
                else:
                    # Digital page
                    # Check for tables using PyMuPDF table finder
                    tables_md = []
                    try:
                        tabs = page.find_tables()
                        for tab in tabs:
                            df_tab = tab.extract()
                            if df_tab:
                                md_tab = self._rows_to_markdown_table(df_tab)
                                tables_md.append(md_tab)
                    except Exception:
                        pass

                    sections.append({
                        "type": "text",
                        "content": raw_text,
                        "page": page_num,
                        "heading": f"Trang {page_num}",
                        "confidence": 1.0,
                    })

                    for tbl in tables_md:
                        sections.append({
                            "type": "table",
                            "content": tbl,
                            "page": page_num,
                            "heading": f"Bảng - Trang {page_num}",
                            "confidence": 1.0,
                        })

        return sections, total_pages, warnings

    # -------------------------------------------------------------------------
    # DOCX Parsing
    # -------------------------------------------------------------------------
    def _parse_docx(self, file_path: str) -> Tuple[List[Dict[str, Any]], int, List[str]]:
        """Parses DOCX document preserving headings, paragraphs, and tables."""
        sections: List[Dict[str, Any]] = []
        warnings: List[str] = []

        doc = docx.Document(file_path)
        current_heading = "Văn bản"

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name.lower() if p.style else ""
            if "heading 1" in style_name:
                current_heading = text
                sections.append({
                    "type": "heading",
                    "level": 1,
                    "content": text,
                    "page": 1,
                    "heading": text,
                    "confidence": 1.0,
                })
            elif "heading 2" in style_name:
                sections.append({
                    "type": "heading",
                    "level": 2,
                    "content": text,
                    "page": 1,
                    "heading": text,
                    "confidence": 1.0,
                })
            elif "heading 3" in style_name:
                sections.append({
                    "type": "heading",
                    "level": 3,
                    "content": text,
                    "page": 1,
                    "heading": text,
                    "confidence": 1.0,
                })
            else:
                sections.append({
                    "type": "text",
                    "content": text,
                    "page": 1,
                    "heading": current_heading,
                    "confidence": 1.0,
                })

        # Extract docx tables as markdown tables
        for tbl in doc.tables:
            grid: List[List[str]] = []
            for row in tbl.rows:
                row_vals = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                grid.append(row_vals)
            if grid:
                md_table = self._rows_to_markdown_table(grid)
                sections.append({
                    "type": "table",
                    "content": md_table,
                    "page": 1,
                    "heading": f"Bảng - {current_heading}",
                    "confidence": 1.0,
                })

        return sections, 1, warnings

    # -------------------------------------------------------------------------
    # Markdown Parsing
    # -------------------------------------------------------------------------
    def _parse_markdown(self, file_path: str) -> Tuple[List[Dict[str, Any]], int, List[str]]:
        """Parses markdown text splitting by headings (#, ##, ###)."""
        sections: List[Dict[str, Any]] = []
        warnings: List[str] = []

        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            lines = f.readlines()

        current_heading = "Tài liệu"
        current_block: List[str] = []

        for line in lines:
            line_str = line.strip()
            heading_match = re.match(r"^(#{1,3})\s+(.*)$", line_str)
            if heading_match:
                if current_block:
                    sections.append({
                        "type": "text",
                        "content": "\n".join(current_block),
                        "page": 1,
                        "heading": current_heading,
                        "confidence": 1.0,
                    })
                    current_block = []

                level = len(heading_match.group(1))
                h_text = heading_match.group(2).strip()
                current_heading = h_text
                sections.append({
                    "type": "heading",
                    "level": level,
                    "content": h_text,
                    "page": 1,
                    "heading": h_text,
                    "confidence": 1.0,
                })
            else:
                if line_str:
                    current_block.append(line_str)

        if current_block:
            sections.append({
                "type": "text",
                "content": "\n".join(current_block),
                "page": 1,
                "heading": current_heading,
                "confidence": 1.0,
            })

        return sections, 1, warnings

    # -------------------------------------------------------------------------
    # Text Parsing
    # -------------------------------------------------------------------------
    def _parse_txt(self, file_path: str) -> Tuple[List[Dict[str, Any]], int, List[str]]:
        """Parses plain text document by paragraphs."""
        sections: List[Dict[str, Any]] = []
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()

        paragraphs = [p.strip() for p in re.split(r"\n\s*\n+", content) if p.strip()]
        for idx, p in enumerate(paragraphs):
            sections.append({
                "type": "text",
                "content": p,
                "page": 1,
                "heading": f"Đoạn {idx + 1}",
                "confidence": 1.0,
            })

        return sections, 1, []

    # -------------------------------------------------------------------------
    # Hierarchical Chunking (Max 512 tokens, 64 overlap, Table splitting)
    # -------------------------------------------------------------------------
    def _hierarchical_chunk(
        self,
        sections: List[Dict[str, Any]],
        doc_id: str,
        file_id: str,
        file_name: str,
        owner_user_id: str,
        department_id: str,
        is_public: bool,
    ) -> List[DocumentChunk]:
        """
        Chunks sections hierarchically:
        - Tracks H1 > H2 breadcrumbs.
        - Text split into max 512 tokens with 64 overlap.
        - Tables kept intact or split by rows repeating markdown header.
        """
        chunks: List[DocumentChunk] = []

        h1 = ""
        h2 = ""
        h3 = ""

        for sec in sections:
            s_type = sec.get("type", "text")
            content = sec.get("content", "").strip()
            page_num = sec.get("page", 1)
            confidence = float(sec.get("confidence", 1.0))

            if not content:
                continue

            # Update heading hierarchy
            if s_type == "heading":
                level = sec.get("level", 1)
                if level == 1:
                    h1 = content
                    h2 = ""
                    h3 = ""
                elif level == 2:
                    h2 = content
                    h3 = ""
                elif level == 3:
                    h3 = content
                continue

            # Build breadcrumb
            path_parts = [file_name]
            if h1:
                path_parts.append(h1)
            if h2:
                path_parts.append(h2)
            if h3:
                path_parts.append(h3)

            breadcrumb = " > ".join(path_parts)
            section_title = h3 or h2 or h1 or sec.get("heading", file_name)
            heading_path = " > ".join([p for p in [h1, h2, h3] if p]) or section_title

            # Table Chunking
            if s_type == "table":
                table_chunks = self._chunk_table(
                    table_content=content,
                    doc_id=doc_id,
                    file_id=file_id,
                    file_name=file_name,
                    breadcrumb=breadcrumb,
                    heading_path=heading_path,
                    section_title=section_title,
                    page_num=page_num,
                    owner_user_id=owner_user_id,
                    department_id=department_id,
                    is_public=is_public,
                    confidence=confidence,
                )
                chunks.extend(table_chunks)
                continue

            # Text Chunking with overlap
            text_tokens = self._tokenize(content)
            if len(text_tokens) <= self.max_chunk_tokens:
                chunks.append(
                    DocumentChunk(
                        chunk_id=str(uuid.uuid4()),
                        doc_id=doc_id,
                        file_id=file_id,
                        file_name=file_name,
                        content=content,
                        breadcrumb=breadcrumb,
                        heading_path=heading_path,
                        section_title=section_title,
                        page_number=page_num,
                        chunk_type="text",
                        owner_user_id=owner_user_id,
                        department_id=department_id,
                        is_public=is_public,
                        ocr_confidence=confidence,
                    )
                )
            else:
                # Sliding window chunking
                step = self.max_chunk_tokens - self.chunk_overlap_tokens
                for start_idx in range(0, len(text_tokens), step):
                    sub_tokens = text_tokens[start_idx : start_idx + self.max_chunk_tokens]
                    sub_text = self._detokenize(sub_tokens)

                    chunks.append(
                        DocumentChunk(
                            chunk_id=str(uuid.uuid4()),
                            doc_id=doc_id,
                            file_id=file_id,
                            file_name=file_name,
                            content=sub_text,
                            breadcrumb=breadcrumb,
                            heading_path=heading_path,
                            section_title=section_title,
                            page_number=page_num,
                            chunk_type="text",
                            owner_user_id=owner_user_id,
                            department_id=department_id,
                            is_public=is_public,
                            ocr_confidence=confidence,
                        )
                    )
                    if start_idx + self.max_chunk_tokens >= len(text_tokens):
                        break

        return chunks

    def _chunk_table(
        self,
        table_content: str,
        doc_id: str,
        file_id: str,
        file_name: str,
        breadcrumb: str,
        heading_path: str,
        section_title: str,
        page_num: int,
        owner_user_id: str,
        department_id: str,
        is_public: bool,
        confidence: float,
    ) -> List[DocumentChunk]:
        """
        Chunks tables. If table exceeds max_chunk_tokens, cuts into row groups
        and repeats the header rows on each chunk.
        """
        tokens = self._tokenize(table_content)
        if len(tokens) <= self.max_chunk_tokens:
            return [
                DocumentChunk(
                    chunk_id=str(uuid.uuid4()),
                    doc_id=doc_id,
                    file_id=file_id,
                    file_name=file_name,
                    content=table_content,
                    breadcrumb=breadcrumb,
                    heading_path=heading_path,
                    section_title=section_title,
                    page_number=page_num,
                    chunk_type="table",
                    owner_user_id=owner_user_id,
                    department_id=department_id,
                    is_public=is_public,
                    ocr_confidence=confidence,
                )
            ]

        # Table exceeds limit -> split by rows and repeat header
        lines = [line.strip() for line in table_content.splitlines() if line.strip()]
        if len(lines) <= 2:
            return [
                DocumentChunk(
                    chunk_id=str(uuid.uuid4()),
                    doc_id=doc_id,
                    file_id=file_id,
                    file_name=file_name,
                    content=table_content,
                    breadcrumb=breadcrumb,
                    heading_path=heading_path,
                    section_title=section_title,
                    page_number=page_num,
                    chunk_type="table",
                    owner_user_id=owner_user_id,
                    department_id=department_id,
                    is_public=is_public,
                    ocr_confidence=confidence,
                )
            ]

        header_lines = lines[:2]  # Header row and delimiter row (|---|---|)
        header_text = "\n".join(header_lines)
        header_tokens = len(self._tokenize(header_text))
        max_rows_tokens = self.max_chunk_tokens - header_tokens - 10

        sub_chunks: List[DocumentChunk] = []
        current_rows: List[str] = []
        current_token_count = 0

        for row in lines[2:]:
            row_tokens = len(self._tokenize(row))
            if current_token_count + row_tokens > max_rows_tokens and current_rows:
                chunk_tbl = header_text + "\n" + "\n".join(current_rows)
                sub_chunks.append(
                    DocumentChunk(
                        chunk_id=str(uuid.uuid4()),
                        doc_id=doc_id,
                        file_id=file_id,
                        file_name=file_name,
                        content=chunk_tbl,
                        breadcrumb=breadcrumb,
                        heading_path=heading_path,
                        section_title=section_title,
                        page_number=page_num,
                        chunk_type="table",
                        owner_user_id=owner_user_id,
                        department_id=department_id,
                        is_public=is_public,
                        ocr_confidence=confidence,
                    )
                )
                current_rows = [row]
                current_token_count = row_tokens
            else:
                current_rows.append(row)
                current_token_count += row_tokens

        if current_rows:
            chunk_tbl = header_text + "\n" + "\n".join(current_rows)
            sub_chunks.append(
                DocumentChunk(
                    chunk_id=str(uuid.uuid4()),
                    doc_id=doc_id,
                    file_id=file_id,
                    file_name=file_name,
                    content=chunk_tbl,
                    breadcrumb=breadcrumb,
                    heading_path=heading_path,
                    section_title=section_title,
                    page_number=page_num,
                    chunk_type="table",
                    owner_user_id=owner_user_id,
                    department_id=department_id,
                    is_public=is_public,
                    ocr_confidence=confidence,
                )
            )

        return sub_chunks

    def _rows_to_markdown_table(self, rows: List[List[Any]]) -> str:
        """Converts 2D list of cells to Markdown table format."""
        if not rows:
            return ""

        clean_grid: List[List[str]] = []
        max_cols = max(len(r) for r in rows)

        for r in rows:
            row_vals = [str(c).strip().replace("\n", " ") if c is not None else "" for c in r]
            if len(row_vals) < max_cols:
                row_vals.extend([""] * (max_cols - len(row_vals)))
            clean_grid.append(row_vals)

        header = "| " + " | ".join(clean_grid[0]) + " |"
        sep = "| " + " | ".join(["---"] * max_cols) + " |"
        data_rows = ["| " + " | ".join(r) + " |" for r in clean_grid[1:]]

        return "\n".join([header, sep] + data_rows)

    def _tokenize(self, text: str) -> List[str]:
        """Tokenizes text using the embedder's tokenizer if available, else regex."""
        if hasattr(self.embedder, "_model") and hasattr(self.embedder._model, "tokenizer"):
            try:
                return self.embedder._model.tokenizer.tokenize(text)
            except Exception:
                pass
        return re.findall(r"\S+|\n", text)

    def _detokenize(self, tokens: List[str]) -> str:
        """Reconstructs text string from tokens."""
        if hasattr(self.embedder, "_model") and hasattr(self.embedder._model, "tokenizer"):
            try:
                tokenizer = self.embedder._model.tokenizer
                token_ids = tokenizer.convert_tokens_to_ids(tokens)
                return tokenizer.decode(token_ids, skip_special_tokens=True)
            except Exception:
                pass
        res = []
        for t in tokens:
            if t == "\n":
                res.append("\n")
            else:
                if res and res[-1] != "\n":
                    res.append(" ")
                res.append(t)
        return "".join(res)
