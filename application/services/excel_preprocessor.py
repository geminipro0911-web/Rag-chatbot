"""
Excel Preprocessor Service.
Handles enterprise-grade ingestion for tabular Excel files (.xlsx).
Follows Module 1 specifications:
- Unmerge rectangular cells (fills top-left value across entire range horizontal & vertical).
- Intelligent multi-tier header detection (combines parent-child as 'Parent - Child').
- Duplicate and empty column sanitization with unidecoded snake_case aliases.
- Summary row detection and stripping ('Tổng', 'Tổng cộng', 'Cộng', 'Total').
- Robust VN vs International number formats, currency, percentages.
- Date parsing (Vietnamese formats, openpyxl datetimes, Excel serial numbers 1990-2100).
- PyArrow ZSTD Parquet storage.
- Entity linking distinct values extraction.
"""

from __future__ import annotations

import os
import re
import uuid
import logging
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import unidecode

logger = logging.getLogger(__name__)


@dataclass
class ColumnMetadata:
    original_name: str
    column_alias: str
    data_type: str
    sample_values: List[Any]
    min_value: Optional[str] = None
    max_value: Optional[str] = None
    null_ratio: float = 0.0


@dataclass
class ProcessedSheet:
    table_id: str
    file_id: str
    source_file: str
    sheet_name: str
    parquet_path: str
    owner_user_id: str
    department_id: str
    is_public: bool
    row_count: int
    columns_schema: List[ColumnMetadata]
    dataframe: pd.DataFrame
    warnings: List[str] = field(default_factory=list)
    distinct_values: Dict[str, List[Dict[str, Any]]] = field(default_factory=dict)


@dataclass
class ExcelProcessResult:
    file_id: str
    file_path: str
    sheets: List[ProcessedSheet]
    all_warnings: List[str] = field(default_factory=list)


class ExcelPreprocessor:
    """Preprocesses enterprise Excel workbooks into normalized parquet tables and metadata."""

    DATE_KEYWORDS = (
        "ngày", "ngay", "date", "thời gian", "thoi_gian", "hạn", "han",
        "kỳ", "ky", "tháng", "thang", "năm", "nam", "period"
    )
    SUMMARY_REGEX = re.compile(
        r"^\s*(tổng|tổng cộng|cộng|tổng số|total|grand total|sum)\b",
        re.IGNORECASE
    )
    VN_DATE_REGEX = re.compile(
        r"ngày\s+(\d{1,2})\s+tháng\s+(\d{1,2})\s+năm\s+(\d{4})",
        re.IGNORECASE
    )

    def __init__(
        self,
        output_dir: str = "data/parquet",
        max_distinct_limit: int = 500,
        distinct_ratio_limit: float = 0.5,
        header_scan_rows: int = 30,
        type_conversion_threshold: float = 0.90,
    ):
        self.output_dir = output_dir
        self.max_distinct_limit = max_distinct_limit
        self.distinct_ratio_limit = distinct_ratio_limit
        self.header_scan_rows = header_scan_rows
        self.type_conversion_threshold = type_conversion_threshold

    def process_file(
        self,
        file_path: str,
        file_id: str,
        department_id: str,
        owner_user_id: str,
        is_public: bool = False,
        skip_hidden: bool = True,
    ) -> ExcelProcessResult:
        """Processes an entire Excel workbook or CSV file and generates parquet files and metadata for each sheet."""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found at {file_path}")

        if file_path.lower().endswith(".csv"):
            return self._process_csv_file(
                file_path=file_path,
                file_id=file_id,
                department_id=department_id,
                owner_user_id=owner_user_id,
                is_public=is_public,
            )

        raw_basename = os.path.basename(file_path)
        source_file = raw_basename[len(file_id) + 1:] if file_id and raw_basename.startswith(f"{file_id}_") else raw_basename
        all_warnings: List[str] = []
        processed_sheets: List[ProcessedSheet] = []

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
        except Exception as exc:
            raise ValueError(f"Failed to read Excel workbook: {exc}") from exc

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheet_warnings: List[str] = []

            # 1. Skip hidden sheet check
            if skip_hidden and ws.sheet_state != "visible":
                msg = f"Sheet '{sheet_name}' is hidden; skipped."
                logger.info(msg)
                all_warnings.append(msg)
                continue

            # 2. Skip empty sheet check
            if ws.max_row is None or ws.max_row == 0 or ws.max_column is None or ws.max_column == 0:
                msg = f"Sheet '{sheet_name}' is completely empty; skipped."
                logger.info(msg)
                all_warnings.append(msg)
                continue

            # Check cached formula values
            self._check_formula_cached_values(ws, sheet_warnings)

            # 3. Unmerge cells (strictly fill top-left across entire rectangular range)
            self._unmerge_and_fill(ws)

            # 4. Extract raw rows
            raw_rows = list(ws.iter_rows(values_only=True))
            if not raw_rows or all(all(cell is None for cell in r) for r in raw_rows):
                msg = f"Sheet '{sheet_name}' contains no data cells; skipped."
                all_warnings.append(msg)
                continue

            # 5. Header detection (with 2-tier support)
            h_start, data_start, raw_headers = self._detect_headers(raw_rows, sheet_name, sheet_warnings)
            if not raw_headers:
                msg = f"Sheet '{sheet_name}' could not identify valid header row; skipped."
                all_warnings.append(msg)
                continue

            # 6. Clean column names & generate aliases
            cleaned_names, column_aliases = self._clean_and_alias_columns(raw_headers)

            # 7. Build DataFrame
            data_rows = raw_rows[data_start:]
            df = self._build_dataframe(data_rows, cleaned_names, column_aliases)
            if df.empty:
                msg = f"Sheet '{sheet_name}' has header but 0 data rows."
                sheet_warnings.append(msg)
                all_warnings.append(msg)
                continue

            # 8. Summary rows detection and stripping
            df, summary_warnings = self._strip_summary_rows(df)
            sheet_warnings.extend(summary_warnings)

            # 9. Type normalization (numbers, currency, percentages, dates, serials)
            df, col_types = self._normalize_column_types(df, sheet_warnings)

            # 10. Distinct values for Entity Linking
            distinct_dict = self._extract_distinct_values(df, col_types)

            # 11. Column catalog metadata
            columns_schema = self._build_column_schema(df, cleaned_names, column_aliases, col_types)

            # 12. Save to Parquet
            table_id = str(uuid.uuid4())
            safe_dept = re.sub(r"[^a-zA-Z0-9_-]", "_", str(department_id))
            dept_parquet_dir = os.path.join(self.output_dir, safe_dept)
            os.makedirs(dept_parquet_dir, exist_ok=True)
            parquet_path = os.path.join(dept_parquet_dir, f"{table_id}.parquet").replace("\\", "/")

            # We write the dataframe with clean aliases as parquet column names for DuckDB ease,
            # but preserve original metadata mappings
            self._save_to_parquet(df, parquet_path)

            processed_sheet = ProcessedSheet(
                table_id=table_id,
                file_id=file_id,
                source_file=source_file,
                sheet_name=sheet_name,
                parquet_path=parquet_path,
                owner_user_id=owner_user_id,
                department_id=department_id,
                is_public=is_public,
                row_count=len(df),
                columns_schema=columns_schema,
                dataframe=df,
                warnings=sheet_warnings,
                distinct_values=distinct_dict,
            )
            processed_sheets.append(processed_sheet)
            all_warnings.extend(sheet_warnings)

        return ExcelProcessResult(
            file_id=file_id,
            file_path=file_path,
            sheets=processed_sheets,
            all_warnings=all_warnings,
        )

    def _process_csv_file(
        self,
        file_path: str,
        file_id: str,
        department_id: str,
        owner_user_id: str,
        is_public: bool = False,
    ) -> ExcelProcessResult:
        """Processes a CSV tabular file and converts to Parquet with full schema extraction."""
        raw_basename = os.path.basename(file_path)
        source_file = raw_basename[len(file_id) + 1:] if file_id and raw_basename.startswith(f"{file_id}_") else raw_basename
        sheet_name = Path(source_file).stem
        encodings = ["utf-8-sig", "utf-8", "cp1258", "latin1"]
        raw_df = None

        for enc in encodings:
            for sep in [",", ";", "\t", "|"]:
                try:
                    sample = pd.read_csv(file_path, encoding=enc, sep=sep, nrows=5)
                    if len(sample.columns) > 1:
                        raw_df = pd.read_csv(file_path, encoding=enc, sep=sep)
                        break
                except Exception:
                    continue
            if raw_df is not None:
                break

        if raw_df is None:
            try:
                raw_df = pd.read_csv(file_path, encoding="utf-8-sig")
            except Exception as e:
                raise ValueError(f"Không thể đọc file CSV '{source_file}': {e}") from e

        # Clean column names & aliases
        clean_cols, aliases = self._clean_and_alias_columns([str(c) for c in raw_df.columns])
        raw_df.columns = aliases

        # Strip summary rows
        df, summary_warnings = self._strip_summary_rows(raw_df)

        # 3. Chuẩn hóa kiểu dữ liệu (số VN, ngày tháng, tiền tệ)
        warnings: List[str] = list(summary_warnings)
        norm_df, col_types = self._normalize_column_types(df, warnings)

        # 4. Trích xuất distinct values cho các cột phân loại
        distinct_vals = self._extract_distinct_values(norm_df, col_types)

        # 5. Xây dựng schema cột
        col_schema = self._build_column_schema(norm_df, clean_cols, aliases, col_types)

        # 6. Save to parquet
        table_id = str(uuid.uuid4())
        dept_parquet_dir = os.path.join(self.output_dir, str(department_id))
        os.makedirs(dept_parquet_dir, exist_ok=True)
        parquet_path = os.path.join(dept_parquet_dir, f"{table_id}.parquet").replace("\\", "/")
        self._save_to_parquet(norm_df, parquet_path)

        processed_sheet = ProcessedSheet(
            table_id=table_id,
            file_id=file_id,
            source_file=source_file,
            sheet_name=sheet_name,
            parquet_path=parquet_path,
            owner_user_id=owner_user_id,
            department_id=department_id,
            is_public=is_public,
            row_count=len(norm_df),
            columns_schema=col_schema,
            dataframe=norm_df,
            warnings=warnings,
            distinct_values=distinct_vals,
        )

        return ExcelProcessResult(
            file_id=file_id,
            file_path=file_path,
            sheets=[processed_sheet],
            all_warnings=warnings,
        )

    def _check_formula_cached_values(self, ws: Any, warnings: List[str]) -> None:
        """Inspects formula cells to warn if cached values are missing."""
        missing_count = 0
        first_coords: List[str] = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" and cell.value is None:
                    missing_count += 1
                    if len(first_coords) < 3:
                        first_coords.append(cell.coordinate)
        if missing_count > 0:
            coords_str = ", ".join(first_coords)
            warnings.append(
                f"Found {missing_count} formula cells without cached values (e.g. {coords_str}). "
                "Open and save file in Microsoft Excel to compute cached results."
            )

    def _unmerge_and_fill(self, ws: Any) -> None:
        """Unmerges all merged ranges and propagates top-left value across the entire rectangular range."""
        merged_ranges = list(ws.merged_cells.ranges)
        for rng in merged_ranges:
            min_col, min_row, max_col, max_row = rng.min_col, rng.min_row, rng.max_col, rng.max_row
            top_left_val = ws.cell(row=min_row, column=min_col).value
            ws.unmerge_cells(str(rng))
            for r in range(min_row, max_row + 1):
                for c in range(min_col, max_col + 1):
                    ws.cell(row=r, column=c).value = top_left_val

    def _detect_headers(
        self,
        raw_rows: List[Tuple[Any, ...]],
        sheet_name: str,
        warnings: List[str],
    ) -> Tuple[int, int, List[str]]:
        """
        Scores first N=30 rows to identify the table header.
        Supports 2-tier headers combined as 'Parent - Child'.
        Returns (header_start_idx, data_start_idx, combined_headers).
        """
        max_scan = min(self.header_scan_rows, len(raw_rows))
        scores: List[Dict[str, Any]] = []

        for idx in range(max_scan):
            row = raw_rows[idx]
            non_null = [c for c in row if c is not None and str(c).strip() != ""]
            if not non_null:
                continue

            non_null_ratio = len(non_null) / max(1, len(row))
            str_items = [c for c in non_null if isinstance(c, str) and not str(c).strip().replace(".", "", 1).isdigit()]
            str_ratio = len(str_items) / len(non_null)

            # Penalize single-cell title lines spanning across wide columns
            is_title = (len(non_null) <= 2 and len(row) > 4)

            # Check next row for subheader candidate
            next_is_subheader = False
            if idx + 1 < len(raw_rows):
                next_row = raw_rows[idx + 1]
                next_non_null = [c for c in next_row if c is not None and str(c).strip() != ""]
                if next_non_null:
                    next_str_items = [c for c in next_non_null if isinstance(c, str)]
                    if (len(next_str_items) / len(next_non_null)) >= 0.6 and len(next_non_null) >= len(non_null) * 0.7:
                        next_is_subheader = True

            score = (non_null_ratio * 0.4) + (str_ratio * 0.6)
            if is_title:
                score *= 0.1

            scores.append({
                "idx": idx,
                "score": score,
                "non_null_ratio": non_null_ratio,
                "str_ratio": str_ratio,
                "next_is_subheader": next_is_subheader,
                "row": row,
            })

        if not scores:
            return -1, -1, []

        scores.sort(key=lambda s: s["score"], reverse=True)
        best = scores[0]
        h_idx = best["idx"]

        # Determine if 2-tier header:
        # A true 2-tier header has horizontal grouping (adjacent identical parent names from unmerged cells)
        # and has data rows below it (h_idx + 2 < len(raw_rows)).
        has_horizontal_grouping = False
        row_h = raw_rows[h_idx]
        for c in range(len(row_h) - 1):
            v1 = row_h[c]
            v2 = row_h[c + 1]
            if v1 is not None and v2 is not None and str(v1).strip() != "" and str(v1).strip() == str(v2).strip():
                has_horizontal_grouping = True
                break

        is_two_tier = False
        if best["next_is_subheader"] and has_horizontal_grouping and (h_idx + 2 < len(raw_rows)):
            is_two_tier = True

        combined_headers: List[str] = []
        num_cols = len(raw_rows[h_idx])

        if is_two_tier:
            row_parent = raw_rows[h_idx]
            row_child = raw_rows[h_idx + 1]
            data_start_idx = h_idx + 2
            warnings.append(f"Sheet '{sheet_name}': Detected 2-tier multi-header at rows {h_idx + 1} and {h_idx + 2}.")

            for c in range(num_cols):
                p_val = str(row_parent[c]).strip() if c < len(row_parent) and row_parent[c] is not None else ""
                c_val = str(row_child[c]).strip() if c < len(row_child) and row_child[c] is not None else ""

                if p_val and c_val and p_val.lower() != c_val.lower():
                    combined = f"{p_val} - {c_val}"
                elif p_val:
                    combined = p_val
                elif c_val:
                    combined = c_val
                else:
                    combined = f"col_{c + 1}"
                combined_headers.append(combined)
        else:
            data_start_idx = h_idx + 1
            for c in range(num_cols):
                val = str(raw_rows[h_idx][c]).strip() if raw_rows[h_idx][c] is not None else ""
                combined_headers.append(val if val else f"col_{c + 1}")

        return h_idx, data_start_idx, combined_headers

    def _clean_and_alias_columns(self, raw_headers: List[str]) -> Tuple[List[str], List[str]]:
        """Cleans original names (strips whitespace/newlines, deduplicates) and creates snake_case aliases."""
        cleaned_names: List[str] = []
        aliases: List[str] = []

        seen_display: Dict[str, int] = {}
        seen_aliases: Dict[str, int] = {}

        for idx, h in enumerate(raw_headers):
            # Clean display name
            cleaned = re.sub(r"[\r\n\t]+", " ", str(h)).strip()
            cleaned = re.sub(r"\s+", " ", cleaned)
            if not cleaned or cleaned.lower().startswith("unnamed"):
                cleaned = f"col_{idx + 1}"

            if cleaned in seen_display:
                seen_display[cleaned] += 1
                unique_display = f"{cleaned}_{seen_display[cleaned]}"
            else:
                seen_display[cleaned] = 1
                unique_display = cleaned
            cleaned_names.append(unique_display)

            # Generate SQL-compatible alias (unidecoded, lowercase snake_case)
            ascii_text = unidecode.unidecode(unique_display).lower()
            alias = re.sub(r"[^a-z0-9]+", "_", ascii_text).strip("_")
            if not alias:
                alias = f"col_{idx + 1}"
            elif alias[0].isdigit():
                alias = f"col_{alias}"

            if alias in seen_aliases:
                seen_aliases[alias] += 1
                unique_alias = f"{alias}_{seen_aliases[alias]}"
            else:
                seen_aliases[alias] = 1
                unique_alias = alias
            aliases.append(unique_alias)

        return cleaned_names, aliases

    def _build_dataframe(
        self,
        data_rows: List[Tuple[Any, ...]],
        cleaned_names: List[str],
        column_aliases: List[str],
    ) -> pd.DataFrame:
        """Constructs pandas DataFrame and removes empty rows/columns."""
        if not data_rows:
            return pd.DataFrame(columns=column_aliases)

        padded_rows = []
        num_cols = len(cleaned_names)
        for r in data_rows:
            row_list = list(r)
            if len(row_list) < num_cols:
                row_list.extend([None] * (num_cols - len(row_list)))
            elif len(row_list) > num_cols:
                row_list = row_list[:num_cols]
            padded_rows.append(row_list)

        df = pd.DataFrame(padded_rows, columns=column_aliases)

        # Drop entirely empty rows
        df = df.dropna(how="all").reset_index(drop=True)

        # Strip whitespace for string cells
        for col in df.columns:
            df[col] = df[col].apply(lambda v: v.strip() if isinstance(v, str) else v)

        return df

    def _strip_summary_rows(self, df: pd.DataFrame) -> Tuple[pd.DataFrame, List[str]]:
        """Identifies and drops summary rows like 'Tổng', 'Tổng cộng', 'Cộng', 'Total'."""
        warnings: List[str] = []
        if df.empty:
            return df, warnings

        summary_indices: List[int] = []
        # Check first 3 columns
        check_cols = list(df.columns[:min(3, len(df.columns))])

        for idx, row in df.iterrows():
            is_summary = False
            summary_label = ""
            for col in check_cols:
                val = str(row[col]) if row[col] is not None else ""
                if val and self.SUMMARY_REGEX.search(val):
                    is_summary = True
                    summary_label = val
                    break
            if is_summary:
                summary_indices.append(idx)
                warnings.append(f"Removed summary row at row index {idx}: '{summary_label}'")

        if summary_indices:
            df = df.drop(index=summary_indices).reset_index(drop=True)

        return df, warnings

    def _normalize_column_types(
        self,
        df: pd.DataFrame,
        warnings: List[str],
    ) -> Tuple[pd.DataFrame, Dict[str, str]]:
        """Normalizes columns into numeric or date types if >= 90% non-null values parse."""
        col_types: Dict[str, str] = {}

        for col in df.columns:
            series = df[col]
            non_null = series.dropna()
            total_valid = len(non_null)

            if total_valid == 0:
                col_types[col] = "string"
                continue

            # 1. Check Date (Excel serial date or date string)
            date_series, is_date = self._try_parse_date_column(col, series, total_valid)
            if is_date:
                df[col] = date_series
                col_types[col] = "date"
                continue

            # 2. Check Numeric (VN vs Intl, currency, percentage)
            num_series, is_num, num_type = self._try_parse_number_column(col, series, total_valid)
            if is_num:
                df[col] = num_series
                col_types[col] = num_type
                continue

            # Default: string
            df[col] = series.astype(str).replace({"None": None, "nan": None, "<NA>": None})
            col_types[col] = "string"

        return df, col_types

    def _try_parse_date_column(
        self,
        col_name: str,
        series: pd.Series,
        total_valid: int,
    ) -> Tuple[Optional[pd.Series], bool]:
        """Tries parsing series into standard datetime/date."""
        col_name_lower = col_name.lower()
        has_date_keyword = any(kw in col_name_lower for kw in self.DATE_KEYWORDS)

        # Check Excel serial numbers (1990-01-01 to 2100-01-01 -> ~ 32874 to 73051)
        serial_candidates = 0
        for val in series.dropna():
            if isinstance(val, (int, float)) and not isinstance(val, bool):
                if 32874 <= val <= 73051:
                    serial_candidates += 1
            elif isinstance(val, str) and val.replace(".", "", 1).isdigit():
                try:
                    num_val = float(val)
                    if 32874 <= num_val <= 73051:
                        serial_candidates += 1
                except ValueError:
                    pass

        # Only interpret raw numbers in [32874, 73051] as Excel serial dates
        # IF the column name explicitly suggests a date/time field.
        # Otherwise regular integer ranges (unit prices, quantities, IDs) would be corrupted.
        if has_date_keyword and (serial_candidates / total_valid) >= 0.5:
            try:
                converted = pd.to_datetime(
                    pd.to_numeric(series, errors="coerce"),
                    unit="D",
                    origin="1899-12-30"
                )
                if (converted.notna().sum() / total_valid) >= 0.5:
                    return converted, True
            except Exception:
                pass

        # Check Date Strings (dd/mm/yyyy, Vietnamese 'ngày D tháng M năm YYYY', etc.)
        parsed_dates: List[Optional[datetime]] = []
        success_count = 0

        for val in series:
            if val is None or pd.isna(val):
                parsed_dates.append(None)
                continue
            if isinstance(val, (datetime, pd.Timestamp)):
                parsed_dates.append(val)
                success_count += 1
                continue

            str_val = str(val).strip()
            # Vietnamese regex: ngày D tháng M năm YYYY
            m = self.VN_DATE_REGEX.search(str_val)
            if m:
                try:
                    day, month, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
                    dt = datetime(year, month, day)
                    parsed_dates.append(dt)
                    success_count += 1
                    continue
                except Exception:
                    pass

            # Dayfirst standard parse (skip dayfirst if already ISO YYYY-MM-DD to avoid warnings)
            try:
                if re.match(r"^\d{4}-\d{2}-\d{2}", str_val):
                    dt = pd.to_datetime(str_val, errors="raise")
                else:
                    dt = pd.to_datetime(str_val, dayfirst=True, errors="raise")
                parsed_dates.append(dt.to_pydatetime() if isinstance(dt, pd.Timestamp) else dt)
                success_count += 1
            except Exception:
                parsed_dates.append(None)

        if (success_count / total_valid) >= self.type_conversion_threshold:
            return pd.to_datetime(parsed_dates), True

        return None, False

    def _try_parse_number_column(
        self,
        col_name: str,
        series: pd.Series,
        total_valid: int,
    ) -> Tuple[Optional[pd.Series], bool, str]:
        """Parses numbers with VN format vs Intl format, currency, and percentages."""
        vn_score = 0
        intl_score = 0

        for val in series.dropna():
            if isinstance(val, (int, float)) and not isinstance(val, bool):
                continue
            s = str(val).strip()
            # Clean currency & percent before analyzing structure
            s_clean = re.sub(r"(?i)[₫đvndvnđ\$usd%]", "", s).strip()
            if not s_clean:
                continue

            if "." in s_clean and "," in s_clean:
                if s_clean.rfind(",") > s_clean.rfind("."):
                    vn_score += 2
                else:
                    intl_score += 2
            elif "," in s_clean:
                parts = s_clean.split(",")
                # If only 1 comma and decimal part is 1, 2, or 4 digits, it's VN decimal (e.g. 97,5 or 106,25)
                if len(parts) == 2 and len(parts[1]) in (1, 2, 4):
                    vn_score += 2
                elif all(len(p) == 3 for p in parts[1:]):
                    intl_score += 2
                else:
                    vn_score += 1
            elif "." in s_clean:
                parts = s_clean.split(".")
                if len(parts) > 2:
                    vn_score += 2
                elif len(parts) == 2 and len(parts[1]) == 3 and len(parts[0]) <= 3:
                    # e.g. 1.500 or 500.000 in VN
                    vn_score += 1
                elif len(parts) == 2 and len(parts[1]) in (1, 2):
                    intl_score += 1

        # Prefer VN format only if there is stronger evidence than standard international/dot notation
        use_vn = vn_score > intl_score

        parsed_values: List[Optional[float]] = []
        is_all_integer = True
        success_count = 0

        for val in series:
            if val is None or pd.isna(val):
                parsed_values.append(None)
                continue
            if isinstance(val, (int, float)) and not isinstance(val, bool):
                num = float(val)
                parsed_values.append(num)
                success_count += 1
                if num % 1 != 0:
                    is_all_integer = False
                continue

            cleaned_str = str(val).strip()
            cleaned_str = re.sub(r"(?i)[₫đvndvnđ\$usd]", "", cleaned_str).strip()

            is_percent = False
            if "%" in cleaned_str:
                is_percent = True
                cleaned_str = cleaned_str.replace("%", "").strip()

            # Normalize separators
            if use_vn:
                # Dot is thousand, comma is decimal
                cleaned_str = cleaned_str.replace(".", "").replace(",", ".")
            else:
                # Comma is thousand, dot is decimal
                cleaned_str = cleaned_str.replace(",", "")

            # Remove spaces
            cleaned_str = cleaned_str.replace(" ", "")

            try:
                num = float(cleaned_str)
                if is_percent:
                    num /= 100.0
                    is_all_integer = False
                parsed_values.append(num)
                success_count += 1
                if num % 1 != 0:
                    is_all_integer = False
            except ValueError:
                parsed_values.append(None)

        if (success_count / total_valid) >= self.type_conversion_threshold:
            res_series = pd.Series(parsed_values, index=series.index)
            num_type = "integer" if is_all_integer else "float"
            return res_series, True, num_type

        return None, False, "string"

    def _extract_distinct_values(
        self,
        df: pd.DataFrame,
        col_types: Dict[str, str],
    ) -> Dict[str, List[Dict[str, Any]]]:
        """Extracts distinct values for Entity Linking on string columns."""
        distinct_dict: Dict[str, List[Dict[str, Any]]] = {}
        total_rows = max(1, len(df))

        for col, c_type in col_types.items():
            if c_type != "string":
                continue

            series = df[col].dropna()
            val_counts = series.value_counts()
            n_unique = len(val_counts)

            # Check threshold n_unique <= 500 AND n_unique / row_count <= 0.5
            if (n_unique <= self.max_distinct_limit) and ((n_unique / total_rows) <= self.distinct_ratio_limit):
                items = val_counts.items()
            else:
                # Over threshold: take top-20 most frequent values
                items = val_counts.head(20).items()

            distinct_list = []
            for orig_val, freq in items:
                orig_str = str(orig_val).strip()
                if not orig_str or orig_str.lower() in ("none", "nan", "<na>"):
                    continue
                norm_str = unidecode.unidecode(orig_str).lower().strip()
                distinct_list.append({
                    "original_value": orig_str,
                    "normalized_value": norm_str,
                    "frequency": int(freq),
                })

            distinct_dict[col] = distinct_list

        return distinct_dict

    def _build_column_schema(
        self,
        df: pd.DataFrame,
        cleaned_names: List[str],
        column_aliases: List[str],
        col_types: Dict[str, str],
    ) -> List[ColumnMetadata]:
        """Builds column metadata schema with sample values and min/max stats."""
        schema: List[ColumnMetadata] = []
        name_map = dict(zip(column_aliases, cleaned_names))

        for alias in column_aliases:
            orig_name = name_map[alias]
            dtype = col_types.get(alias, "string")
            series = df[alias].dropna()

            # 3 sample values
            samples = [v for v in series.head(3).tolist() if v is not None]

            min_val = None
            max_val = None
            if dtype in ("integer", "float", "date") and not series.empty:
                try:
                    min_val = str(series.min())
                    max_val = str(series.max())
                except Exception:
                    pass

            null_count = df[alias].isna().sum()
            null_ratio = float(null_count / max(1, len(df)))

            schema.append(
                ColumnMetadata(
                    original_name=orig_name,
                    column_alias=alias,
                    data_type=dtype,
                    sample_values=samples,
                    min_value=min_val,
                    max_value=max_val,
                    null_ratio=round(null_ratio, 4),
                )
            )

        return schema

    def _save_to_parquet(self, df: pd.DataFrame, parquet_path: str) -> None:
        """Saves dataframe using PyArrow table with ZSTD compression."""
        table = pa.Table.from_pandas(df)
        pq.write_table(table, parquet_path, compression="zstd")
        logger.info(f"Saved Parquet table ({len(df)} rows) to {parquet_path}")
